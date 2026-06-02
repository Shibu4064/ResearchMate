from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".csv"}


@dataclass
class LoadedDocument:
    text: str
    source: str
    doc_type: str
    page: int | None = None


def _clean_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def load_pdf(path: Path) -> list[LoadedDocument]:
    docs: list[LoadedDocument] = []
    with pymupdf.open(path) as pdf:
        for idx, page in enumerate(pdf, start=1):
            text = _clean_text(page.get_text("text"))
            if text:
                docs.append(
                    LoadedDocument(
                        text=text,
                        source=path.name,
                        doc_type="pdf",
                        page=idx,
                    )
                )
    return docs


def load_docx(path: Path) -> list[LoadedDocument]:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = _clean_text("\n".join(paragraphs))
    return [LoadedDocument(text=text, source=path.name, doc_type="docx")]


def load_text(path: Path) -> list[LoadedDocument]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [LoadedDocument(text=_clean_text(text), source=path.name, doc_type=path.suffix.lower().lstrip("."))]


def load_csv(path: Path, max_rows: int = 3000) -> list[LoadedDocument]:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            if i >= max_rows:
                rows.append(f"[Truncated after {max_rows} rows]")
                break
            rows.append(" | ".join(f"{k}: {v}" for k, v in row.items()))
    text = _clean_text("\n".join(rows))
    return [LoadedDocument(text=text, source=path.name, doc_type="csv")]


def load_document(path: str | Path) -> list[LoadedDocument]:
    path = Path(path)
    ext = path.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {sorted(SUPPORTED_EXTENSIONS)}")

    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in {".txt", ".md"}:
        return load_text(path)
    if ext == ".csv":
        return load_csv(path)

    raise ValueError(f"Unsupported file type: {ext}")


def load_many(paths: Iterable[str | Path]) -> list[LoadedDocument]:
    all_docs: list[LoadedDocument] = []
    for path in paths:
        all_docs.extend(load_document(path))
    return all_docs
